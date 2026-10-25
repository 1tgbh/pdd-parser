"""
数据导出模块 - 支持 JSON/CSV/Excel/TXT 格式导出

TXT 导出包含商品信息 + 实际图片（多线程并行下载），可直接用于上货。
"""

import csv
import json
import os
import random
import sqlite3
import tempfile
import time
import zipfile
import requests as _req
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path

import urllib3
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

DATA_DIR: Path = Path(__file__).resolve().parent / "data"
EXPORT_DIR: Path = DATA_DIR / "exports"
DB_PATH: Path = DATA_DIR / "goods.db"

EXPORT_DIR.mkdir(parents=True, exist_ok=True)

# ---- 导出时自动改价支持 ----

def _get_export_markup():
    """读取 config.json 中的 price_markup 配置"""
    try:
        cfg_path = DATA_DIR / "config.json"
        if cfg_path.exists():
            with open(cfg_path, "r", encoding="utf-8") as f:
                cfg = json.load(f)
            markup = cfg.get("price_markup", {})
            if markup.get("enabled"):
                return markup
    except Exception:
        pass
    return None


def _apply_export_markup(price, markup):
    """对单个价格应用改价公式"""
    if not markup or price <= 0:
        return price
    return round(price * markup.get("multiplier", 1) + markup.get("fixed_amount", 0), 2)


def _apply_markup_to_goods(goods, markup):
    """对一个商品字典的所有价格字段应用改价"""
    for field in ("price", "original_price", "min_group_price", "price_min", "price_max"):
        v = goods.get(field, 0) or 0
        if v > 0:
            goods[field] = _apply_export_markup(v, markup)
    # SKU 价格
    for sku in goods.get("skus", []):
        if isinstance(sku, dict):
            if (sku.get("price") or 0) > 0:
                sku["price"] = _apply_export_markup(sku["price"], markup)
            if (sku.get("original_price") or 0) > 0:
                sku["original_price"] = _apply_export_markup(sku["original_price"], markup)


_IMG_HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/125.0.0.0 Safari/537.36",
    "Referer": "https://mobile.yangkeduo.com/",
    "Accept": "image/webp,image/apng,image/*,*/*;q=0.8",
}
_IMG_SESSION: _req.Session | None = None


def _get_img_session() -> _req.Session:
    global _IMG_SESSION
    if _IMG_SESSION is None:
        _IMG_SESSION = _req.Session()
        _IMG_SESSION.headers.update(_IMG_HEADERS)
    return _IMG_SESSION


def _download_image(task: tuple) -> tuple:
    """下载单张图片。task = (arcname, url)，返回 (arcname, bytes|None)"""
    arcname, url = task
    if not url or not url.startswith("http"):
        return (arcname, None)
    time.sleep(random.uniform(0.2, 0.8))  # 随机间隔，降低请求突发
    sess = _get_img_session()
    for attempt in range(3):
        try:
            resp = sess.get(url, timeout=15, verify=False)
            if resp.status_code == 200 and len(resp.content) > 200:
                return (arcname, resp.content)
        except Exception:
            if attempt < 2:
                time.sleep(0.3 * (attempt + 1))
    return (arcname, None)


def get_conn() -> sqlite3.Connection:
    """获取数据库连接（Row 模式）。"""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def _sku_specs_to_dict(specs) -> dict:
    """将 SKU specs 统一转为 {key: value} 字典格式。
    
    数据库中 specs 可能是:
      - dict: {"颜色": "红"} (旧格式)
      - list: [{"key": "颜色", "value": "红"}] (新格式)
    """
    if not specs:
        return {}
    if isinstance(specs, dict):
        return specs
    if isinstance(specs, list):
        result = {}
        for item in specs:
            if isinstance(item, dict):
                k = item.get("key", item.get("spec_key", ""))
                v = item.get("value", item.get("spec_value", ""))
                if k:
                    result[k] = v
        return result
    return {}


def _load_goods(rows) -> list:
    """将数据库行转为商品字典列表"""
    goods_list = []
    for r in rows:
        goods = {
            "goods_id": r["goods_id"],
            "title": r["title"] or "",
            "subtitle": r["subtitle"] or "",
            "selling_points": json.loads(r["selling_points"] or "[]"),
            "price": r["price"] or 0,
            "original_price": r["original_price"] or 0,
            "min_group_price": r["min_group_price"] or 0,
            "sales": r["sales"] or 0,
            "sold_quantity": r["sold_quantity"] or 0,
            "shop_name": r["shop_name"] or "",
            "shop_id": r["shop_id"] or "",
            "main_images": json.loads(r["main_images"] or "[]"),
            "detail_images": json.loads(r["detail_images"] or "[]"),
            "sku_images": json.loads(r["sku_images"] or "[]"),
            "specs": json.loads(r["specs"] or "{}"),
            "skus": json.loads(r["skus"] or "[]"),
            "description": r["description"] or "",
            "attributes": json.loads(r["attributes"] or "{}"),
            "source_url": r["source_url"] or "",
            "crawl_time": r["crawl_time"] or "",
            "price_min": r["price_min"] if "price_min" in r.keys() else (r["price"] or 0),
            "price_max": r["price_max"] if "price_max" in r.keys() else (r["price"] or 0),
            "total_stock": r["total_stock"] if "total_stock" in r.keys() else 0,
        }
        if not goods["skus"]:
            try:
                conn2 = sqlite3.connect(str(DATA_DIR / "goods.db"))
                conn2.row_factory = sqlite3.Row
                sku_rows = conn2.execute("SELECT * FROM sku_items WHERE goods_id = ?", (r["goods_id"],)).fetchall()
                conn2.close()
                for sr in sku_rows:
                    specs_dict = json.loads(sr["specs"] or "{}")
                    goods["skus"].append({
                        "sku_id": sr["sku_id"] or "",
                        "specs": specs_dict,
                        "price": sr["price"] or 0,
                        "original_price": sr["original_price"] or 0,
                        "stock": sr["stock"] or 0,
                        "image": sr["image"] or "",
                    })
            except Exception:
                pass
        goods_list.append(goods)
    # 导出时应用改价公式（AOP自动改价）
    _markup = _get_export_markup()
    if _markup:
        for _g in goods_list:
            _apply_markup_to_goods(_g, _markup)
    return goods_list


def _safe_filename(s: str, max_len: int = 30) -> str:
    bad = '/\\:*?"<>|'
    for ch in bad:
        s = s.replace(ch, "_")
    return s[:max_len]

def generate_goods_txt(g: dict) -> str:
    """生成单个商品的 TXT 描述文本（可复用于导出与打包）。"""
    gid = g["goods_id"]
    title = g["title"] or "unknown"
    lines = [
        "=" * 50,
        f"  {title}",
        "=" * 50, "",
        f"商品ID: {gid}",
        f"商品名称: {title}",
    ]
    if g.get("subtitle"):
        lines.append(f"副标题: {g['subtitle']}")
    lines.append(f"价格: {g['price']:.2f} 元")
    if g.get("original_price"):
        lines.append(f"原价: {g['original_price']:.2f} 元")
    if g.get("min_group_price"):
        lines.append(f"拼团价: {g['min_group_price']:.2f} 元")
    if g.get("sales"):
        lines.append(f"销量: {g['sales']} 件")
    if g.get("sold_quantity"):
        lines.append(f"已拼件数: {g['sold_quantity']} 件")
    if g.get("shop_name"):
        lines.append(f"店铺名称: {g['shop_name']}")
    if g.get("shop_id"):
        lines.append(f"店铺ID: {g['shop_id']}")
    lines.append("")
    if g.get("selling_points"):
        lines.append("卖点:")
        for sp in g["selling_points"]:
            lines.append(f"  - {sp}")
        lines.append("")
    if g.get("specs"):
        lines.append("规格参数:")
        for k, vals in g["specs"].items():
            lines.append(f"  {k}: {', '.join(vals)}")
        lines.append("")
    if g.get("skus"):
        lines.append("SKU详情:")
        for sk in g["skus"]:
            spec_str = " ".join(_sku_specs_to_dict(sk.get("specs")).values()) if sk.get("specs") else ""
            price_str = f"  {spec_str} - {sk['price']:.2f}元" if spec_str else f"  {sk['price']:.2f}元"
            if sk.get("stock"):
                price_str += f" - 库存{sk['stock']}"
            lines.append(price_str)
        lines.append("")
    if g.get("attributes"):
        lines.append("商品属性:")
        for k, v in g["attributes"].items():
            lines.append(f"  {k}: {v}")
        lines.append("")
    if g.get("description"):
        lines.append("商品描述:")
        lines.append(g["description"][:1000])
        lines.append("")
    lines.append(f"商品链接: {g.get('source_url', '')}")
    lines.append(f"采集时间: {g.get('crawl_time', '')}")
    lines.append("")
    if g.get("main_images"):
        lines.append("主图:")
        for i in range(len(g["main_images"])):
            lines.append(f"  {i+1}. 主图_{i+1}.jpg")
        lines.append("")
    if g.get("detail_images"):
        lines.append("详情图:")
        for i in range(len(g["detail_images"])):
            lines.append(f"  {i+1}. 详情图_{i+1}.jpg")
    return "\n".join(lines)
# ==================== 导出函数 ====================

def export_json(goods_ids: list = None) -> str:
    """导出为 JSON 格式"""
    conn = get_conn()
    if goods_ids:
        placeholders = ",".join(["?"] * len(goods_ids))
        rows = conn.execute(f"SELECT * FROM goods WHERE goods_id IN ({placeholders})", goods_ids).fetchall()
    else:
        rows = conn.execute("SELECT * FROM goods ORDER BY crawl_time DESC").fetchall()
    conn.close()
    goods_list = _load_goods(rows)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filepath = EXPORT_DIR / f"goods_{timestamp}.json"
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(goods_list, f, ensure_ascii=False, indent=2)
    print(f"JSON export done: {filepath} ({len(goods_list)} goods)")
    return filepath


def export_csv(goods_ids: list = None) -> str:
    """导出为 CSV 格式"""
    conn = get_conn()
    if goods_ids:
        placeholders = ",".join(["?"] * len(goods_ids))
        rows = conn.execute(f"SELECT * FROM goods WHERE goods_id IN ({placeholders})", goods_ids).fetchall()
    else:
        rows = conn.execute("SELECT * FROM goods ORDER BY crawl_time DESC").fetchall()
    conn.close()
    goods_list = _load_goods(rows)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filepath = EXPORT_DIR / f"goods_{timestamp}.csv"
    headers = [
        "商品ID", "标题", "副标题", "卖点", "价格", "原价", "拼团价",
        "销量", "已拼件数", "店铺名", "规格", "SKU详情",
        "主图链接", "详情图链接", "商品描述", "商品属性"
    ]
    with open(filepath, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for g in goods_list:
            specs_str = "; ".join([f"{k}: {','.join(v)}" for k, v in g["specs"].items()])
            sku_parts = []
            for s in g["skus"]:
                spec_str = " ".join([f"{k}:{v}" for k, v in _sku_specs_to_dict(s.get("specs")).items()])
                sku_parts.append(f"{spec_str} ¥{s.get('price', 0)} 库存{s.get('stock', 0)}")
            skus_str = " | ".join(sku_parts)
            writer.writerow([
                g["goods_id"], g["title"], g["subtitle"],
                "、".join(g["selling_points"]),
                g["price"], g["original_price"], g["min_group_price"],
                g["sales"], g["sold_quantity"], g["shop_name"],
                specs_str, skus_str,
                " ".join(g["main_images"]), " ".join(g["detail_images"]),
                g["description"][:500],
                "; ".join([f"{k}:{v}" for k, v in g["attributes"].items()]),
            ])
    print(f"CSV export done: {filepath} ({len(goods_list)} goods)")
    return filepath


def export_excel(goods_ids: list = None) -> str:
    """导出为 Excel 格式"""
    try:
        import openpyxl
    except ImportError:
        print("需要安装 openpyxl: pip install openpyxl")
        return None
    conn = get_conn()
    if goods_ids:
        placeholders = ",".join(["?"] * len(goods_ids))
        rows = conn.execute(f"SELECT * FROM goods WHERE goods_id IN ({placeholders})", goods_ids).fetchall()
    else:
        rows = conn.execute("SELECT * FROM goods ORDER BY crawl_time DESC").fetchall()
    conn.close()
    goods_list = _load_goods(rows)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filepath = EXPORT_DIR / f"goods_{timestamp}.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "商品数据"
    headers = [
        "商品ID", "标题", "副标题", "卖点", "价格", "原价", "拼团价",
        "销量", "已拼件数", "店铺名", "规格", "SKU详情",
        "主图链接", "详情图链接", "商品描述", "商品属性"
    ]
    ws.append(headers)
    for g in goods_list:
        specs_str = "; ".join([f"{k}: {','.join(v)}" for k, v in g["specs"].items()])
        sku_parts = []
        for s in g["skus"]:
            spec_str = " ".join([f"{k}:{v}" for k, v in _sku_specs_to_dict(s.get("specs")).items()])
            sku_parts.append(f"{spec_str} ¥{s.get('price', 0)} 库存{s.get('stock', 0)}")
        skus_str = " | ".join(sku_parts)
        ws.append([
            g["goods_id"], g["title"], g["subtitle"],
            "、".join(g["selling_points"]),
            g["price"], g["original_price"], g["min_group_price"],
            g["sales"], g["sold_quantity"], g["shop_name"],
            specs_str, skus_str,
            " ".join(g["main_images"]), " ".join(g["detail_images"]),
            g["description"][:500],
            "; ".join([f"{k}:{v}" for k, v in g["attributes"].items()]),
        ])
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_length + 2, 50)
    wb.save(filepath)
    print(f"Excel export done: {filepath} ({len(goods_list)} goods)")
    return filepath


def export_with_template(goods_list: list, fields: list, column_map: dict) -> tuple:
    try:
        import openpyxl
    except ImportError:
        print("需要安装 openpyxl: pip install openpyxl")
        return None, "export.xlsx"
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "商品数据"
    headers = [column_map.get(f, f) for f in fields]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True)
    for row_idx, g in enumerate(goods_list, 2):
        for col_idx, field in enumerate(fields, 1):
            value = getattr(g, field, None) if hasattr(g, 'get') is False else g.get(field, "")
            if isinstance(value, list):
                value = ", ".join(str(v) for v in value)
            elif isinstance(value, dict):
                value = json.dumps(value, ensure_ascii=False)
            ws.cell(row=row_idx, column=col_idx, value=value)
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False, mode="wb")
    wb.save(tmp.name)
    tmp.close()
    return tmp.name, "export.xlsx"


def export_for_taobao(goods_ids: list = None) -> str:
    """导出淘宝/1688上货格式"""
    conn = get_conn()
    if goods_ids:
        placeholders = ",".join(["?"] * len(goods_ids))
        rows = conn.execute(f"SELECT * FROM goods WHERE goods_id IN ({placeholders})", goods_ids).fetchall()
    else:
        rows = conn.execute("SELECT * FROM goods ORDER BY crawl_time DESC").fetchall()
    conn.close()
    goods_list = _load_goods(rows)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filepath = EXPORT_DIR / f"taobao_{timestamp}.json"
    taobao_goods = []
    for g in goods_list:
        item = {
            "outer_id": g["goods_id"],
            "title": g["title"],
            "sub_title": g["subtitle"],
            "price": str(g["price"]),
            "stuff_status": "1",
            "description": g["description"] or g["title"],
            "seller_cids": "", "props_name": "", "input_pids": "", "input_str": "",
            "sku_properties": "", "sku_quantities": "", "sku_prices": "", "sku_outer_ids": "",
            "picture": g["main_images"][0] if g["main_images"] else "",
            "pictures": ",".join(g["main_images"][:5]),
            "location": "",
            "item_weight": g["attributes"].get("重量", ""),
        }
        if g["skus"]:
            sku_props, sku_quantities, sku_prices = [], [], []
            for s in g["skus"]:
                props = ";".join([f"{k}:{v}" for k, v in _sku_specs_to_dict(s.get("specs")).items()])
                sku_props.append(props)
                sku_quantities.append(str(s.get("stock", 0)))
                sku_prices.append(str(s.get("price", g["price"])))
            item["sku_properties"] = "|".join(sku_props)
            item["sku_quantities"] = "|".join(sku_quantities)
            item["sku_prices"] = "|".join(sku_prices)
        taobao_goods.append(item)
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(taobao_goods, f, ensure_ascii=False, indent=2)
    print(f"Taobao export done: {filepath} ({len(taobao_goods)} goods)")
    return filepath


def export_txt(goods_ids: list = None) -> str:
    """导出为 TXT 格式（多线程并行下载图片）- 打包为 ZIP"""
    conn = get_conn()
    if goods_ids:
        placeholders = ",".join(["?"] * len(goods_ids))
        rows = conn.execute(f"SELECT * FROM goods WHERE goods_id IN ({placeholders})", goods_ids).fetchall()
    else:
        rows = conn.execute("SELECT * FROM goods ORDER BY crawl_time DESC").fetchall()
    conn.close()
    goods_list = _load_goods(rows)
    if not goods_list:
        return None
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    zip_path = EXPORT_DIR / f"goods_txt_{timestamp}.zip"
    goods_data = []
    all_tasks = []
    goods_ai_files = []
    for g in goods_list:
        gid = g["goods_id"]
        title = g["title"] or "unknown"
        base_dir = f"{gid}_{_safe_filename(title)}"
        txt_content = generate_goods_txt(g)
        img_tasks = []
        if g["main_images"]:
            for i, url in enumerate(g["main_images"]):
                img_tasks.append((f"{base_dir}/主图/主图_{i+1}.jpg", url))
        if g["detail_images"]:
            for i, url in enumerate(g["detail_images"]):
                img_tasks.append((f"{base_dir}/详情图/详情图_{i+1}.jpg", url))
        # Collect AI-generated images from disk
        ai_conn = get_conn()
        ai_rows = ai_conn.execute(
            "SELECT result_path, prompt FROM ai_images WHERE goods_id = ?", (gid,)
        ).fetchall()
        ai_conn.close()
        ai_local = []
        for ai_idx, ai_row in enumerate(ai_rows, 1):
            ai_path = DATA_DIR / ai_row[0]
            if ai_path.exists():
                ext = ai_path.suffix or ".png"
                ai_local.append((f"{base_dir}/AI图_{ai_idx}{ext}", str(ai_path)))

        goods_data.append((f"{base_dir}/商品信息.txt", txt_content))
        all_tasks.extend(img_tasks)
        goods_ai_files.extend(ai_local)
    results = {}
    total_imgs = len(all_tasks)
    if total_imgs > 0:
        workers = min(3, total_imgs)  # 低并发 + 随机间隔，降低风控风险
        print(f"Downloading {total_imgs} images with {workers} threads...")
        with ThreadPoolExecutor(max_workers=workers) as pool:
            futures = {pool.submit(_download_image, task): task for task in all_tasks}
            for future in as_completed(futures):
                arcname, data = future.result()
                if data:
                    results[arcname] = data
    with zipfile.ZipFile(str(zip_path), "w", zipfile.ZIP_DEFLATED) as zf:
        for txt_arcname, txt_content in goods_data:
            zf.writestr(txt_arcname, txt_content.encode("utf-8-sig"))
        for arcname, data in results.items():
            zf.writestr(arcname, data)
        # Write AI-generated images from local disk
        for arcname, local_path in goods_ai_files:
            zf.write(local_path, arcname)
    downloaded = len(results)
    failed = total_imgs - downloaded
    print(f"TXT export done: {zip_path}")
    print(f"  {len(goods_list)} goods, {downloaded} images OK, {failed} failed")
    return zip_path
def export_dawang(goods_ids: list = None, template: str = "universal") -> str:
    """导出大王解析兼容格式（CSV）

    template:
    - "universal": 通用数据包格式
    - "wechat": 微信小店上货格式
    - "pdd": 拼多多上货格式
    """
    conn = get_conn()
    if goods_ids:
        placeholders = ",".join(["?"] * len(goods_ids))
        rows = conn.execute(
            f"SELECT * FROM goods WHERE goods_id IN ({placeholders})", goods_ids
        ).fetchall()
    else:
        rows = conn.execute("SELECT * FROM goods ORDER BY crawl_time DESC").fetchall()
    conn.close()
    goods_list = _load_goods(rows)
    if not goods_list:
        return None
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    if template == "wechat":
        return _export_dawang_wechat(goods_list, timestamp)
    elif template == "pdd":
        return _export_dawang_pdd(goods_list, timestamp)
    else:
        return _export_dawang_universal(goods_list, timestamp)


def _export_dawang_universal(goods_list: list, timestamp: str) -> str:
    filepath = EXPORT_DIR / f"dawang_universal_{timestamp}.csv"
    headers = [
        "商品ID", "商品标题", "副标题", "当前价格", "原价", "拼团价",
        "销量", "已拼件数", "店铺名称", "店铺ID",
        "主图链接", "详情图链接", "商品描述",
        "规格信息", "SKU详情", "商品属性", "商品链接", "采集时间",
    ]
    with open(filepath, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for g in goods_list:
            specs_str = "; ".join(
                [f"{k}:{','.join(v)}" for k, v in g["specs"].items()]
            )
            sku_parts = []
            for s in g["skus"]:
                spec_vals = " ".join(_sku_specs_to_dict(s.get("specs")).values()) if s.get("specs") else ""
                price_str = f"{s.get('price', 0):.2f}"
                stock_str = str(s.get("stock", 999))
                sku_parts.append(f"{spec_vals}/{price_str}/{stock_str}")
            skus_str = "|".join(sku_parts)
            attrs_str = "; ".join(
                [f"{k}:{v}" for k, v in g["attributes"].items()]
            )
            writer.writerow([
                g["goods_id"], g["title"], g["subtitle"],
                f"{g['price']:.2f}", f"{g['original_price']:.2f}",
                f"{g['min_group_price']:.2f}",
                g["sales"], g["sold_quantity"],
                g["shop_name"], g["shop_id"],
                "|".join(g["main_images"]),
                "|".join(g["detail_images"]),
                g["description"][:2000],
                specs_str, skus_str, attrs_str,
                g["source_url"], g.get("crawl_time", ""),
            ])
    print(f"Dawang universal export: {filepath} ({len(goods_list)} goods)")
    return filepath


def _export_dawang_wechat(goods_list: list, timestamp: str) -> str:
    filepath = EXPORT_DIR / f"dawang_wechat_{timestamp}.csv"
    max_main = 9
    max_detail = 5
    headers = ["商品标题", "副标题", "价格", "原价", "销量", "库存", "商品描述"]
    for i in range(1, max_main + 1):
        headers.append(f"主图{i}")
    for i in range(1, max_detail + 1):
        headers.append(f"详情图{i}")
    headers.extend([
        "规格名1", "规格值1", "规格名2", "规格值2",
        "SKU规格", "SKU价格", "SKU库存", "SKU图片",
        "品牌", "发货方式", "七天无理由",
    ])
    with open(filepath, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for g in goods_list:
            title = g["title"][:60] if g["title"] else ""
            row = [
                title, g["subtitle"][:20] if g["subtitle"] else "",
                f"{g['price']:.2f}", f"{g['original_price']:.2f}",
                g["sales"], 999, g["description"][:2000],
            ]
            main_imgs = g["main_images"][:max_main]
            for i in range(max_main):
                row.append(main_imgs[i] if i < len(main_imgs) else "")
            detail_imgs = g["detail_images"][:max_detail]
            for i in range(max_detail):
                row.append(detail_imgs[i] if i < len(detail_imgs) else "")
            spec_names = list(g["specs"].keys())
            spec_vals_list = list(g["specs"].values())
            row.append(spec_names[0] if len(spec_names) > 0 else "")
            row.append(",".join(spec_vals_list[0]) if len(spec_vals_list) > 0 else "")
            row.append(spec_names[1] if len(spec_names) > 1 else "")
            row.append(",".join(spec_vals_list[1]) if len(spec_vals_list) > 1 else "")
            sku_specs_parts = []
            sku_prices = []
            sku_stocks = []
            sku_imgs = []
            for s in g["skus"]:
                spec_str = " ".join(_sku_specs_to_dict(s.get("specs")).values()) if s.get("specs") else ""
                sku_specs_parts.append(spec_str)
                sku_prices.append(f"{s.get('price', g['price']):.2f}")
                sku_stocks.append(str(s.get("stock", 999)))
                sku_imgs.append(s.get("image", ""))
            row.append("|".join(sku_specs_parts))
            row.append("|".join(sku_prices))
            row.append("|".join(sku_stocks))
            row.append("|".join(sku_imgs))
            brand = g["attributes"].get("品牌", "") if g.get("attributes") else ""
            row.extend([brand, "快递发货", "是"])
            writer.writerow(row)
    print(f"Dawang WeChat export: {filepath} ({len(goods_list)} goods)")
    return filepath


def _export_dawang_pdd(goods_list: list, timestamp: str) -> str:
    filepath = EXPORT_DIR / f"数据包_拼多多上货_{timestamp}.csv"
    headers = [
        "商品ID", "商品标题", "副标题", "当前价格", "原价",
        "销量", "库存", "店铺名称", "商品描述",
        "主图链接", "详情图链接",
        "规格名称", "规格值", "SKU价格", "SKU库存",
        "商品属性", "商品链接",
    ]
    with open(filepath, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for g in goods_list:
            spec_names = "; ".join(g["specs"].keys())
            spec_vals = "; ".join(
                [",".join(v) for v in g["specs"].values()]
            )
            sku_prices = "|".join(
                [f"{s.get('price', 0):.2f}" for s in g["skus"]]
            )
            sku_stocks = "|".join(
                [str(s.get("stock", 999)) for s in g["skus"]]
            )
            total_stock = sum(s.get("stock", 999) for s in g["skus"]) or 999
            attrs_str = "; ".join(
                [f"{k}:{v}" for k, v in g["attributes"].items()]
            )
            writer.writerow([
                g["goods_id"], g["title"], g["subtitle"],
                f"{g['price']:.2f}", f"{g['original_price']:.2f}",
                g["sales"], total_stock,
                g["shop_name"], g["description"][:2000],
                " ".join(g["main_images"]),
                " ".join(g["detail_images"]),
                spec_names, spec_vals,
                sku_prices, sku_stocks,
                attrs_str, g["source_url"],
            ])
    print(f"Dawang PDD export: {filepath} ({len(goods_list)} goods)")
    return filepath


def _normalize_to_dawang_format(data: dict) -> dict:
    """将不同来源的 raw_data 统一转为大王解析格式: {"ts":..., "store": {"initDataObj":...}}
    
    DB中存在3种格式:
    1. 大王解析导入: {"ts":..., "store": {"initDataObj": {...}}} - 已是标准格式
    2. 系统采集(initDataObj): {"goods":{...},"mall":{...},...} - 需要包装
    3. 旧格式(简化): {"goods_id":...,"attributes":...} - 无法还原完整数据
    """
    # 格式1: 已经是大王解析标准格式
    if "store" in data and isinstance(data.get("store"), dict):
        if "initDataObj" in data["store"]:
            return data
    # 格式2: initDataObj扁平结构 (有goods字段)
    if "goods" in data and isinstance(data.get("goods"), dict):
        ts = int(data.get("ts", data.get("goods", {}).get("eventTime", 0)))
        if ts < 1000000000000:
            ts = ts * 1000
        return {"ts": ts, "store": {"initDataObj": data}}
    # 格式3: 旧简化格式 - 包装为最小可用结构
    goods_id = data.get("goods_id", "")
    mall_info = data.get("mall_info") or {}
    return {
        "ts": 0,
        "store": {
            "initDataObj": {
                "goods": {
                    "goodsID": goods_id,
                    "goodsName": data.get("title", ""),
                    "topGallery": [{"url": u} for u in data.get("main_images", [])],
                    "detailGallery": [{"url": u} for u in data.get("detail_images", [])],
                    "skus": [],
                    "goodsProperty": [
                        {"key": k, "values": [v]} for k, v in data.get("attributes", {}).items()
                    ],
                },
                "mall": {
                    "mallName": mall_info.get("mall_name", data.get("shop_name", "")),
                    "mallID": str(mall_info.get("mall_id", data.get("shop_id", ""))),
                },
            }
        },
    }


def export_dawang_raw(goods_ids: list = None) -> str:
    conn = get_conn()
    if goods_ids:
        placeholders = ",".join(["?"] * len(goods_ids))
        rows = conn.execute(
            f"SELECT goods_id, raw_data FROM goods WHERE goods_id IN ({placeholders})", goods_ids
        ).fetchall()
    else:
        rows = conn.execute("SELECT goods_id, raw_data FROM goods ORDER BY crawl_time DESC").fetchall()
    conn.close()
    if not rows:
        return None
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir = EXPORT_DIR / f"dawang_raw_{timestamp}"
    out_dir.mkdir(parents=True, exist_ok=True)
    count = 0
    for r in rows:
        raw = r["raw_data"]
        if not raw:
            continue
        try:
            data = json.loads(raw)
        except Exception:
            continue
        normalized = _normalize_to_dawang_format(data)
        filepath = out_dir / f"{r['goods_id']}.txt"
        with open(filepath, "w", encoding="utf-8") as f:
            json.dump(normalized, f, ensure_ascii=False)
        count += 1
    if count == 0:
        return None
    zip_path = EXPORT_DIR / f"dawang_raw_{timestamp}.zip"
    with zipfile.ZipFile(str(zip_path), "w", zipfile.ZIP_DEFLATED) as zf:
        for f in out_dir.iterdir():
            zf.write(str(f), f.name)
    import shutil
    shutil.rmtree(str(out_dir), ignore_errors=True)
    print(f"Dawang raw export: {zip_path} ({count} goods)")
    return zip_path


if __name__ == "__main__":
    print("=" * 50)
    print("  数据导出工具")
    print("=" * 50)
    if not DB_PATH.exists():
        print("\n数据库不存在，请先采集数据")
        exit(1)
    conn = get_conn()
    count = conn.execute("SELECT COUNT(*) FROM goods").fetchone()[0]
    conn.close()
    print(f"\n数据库中有 {count} 个商品")
    if count == 0:
        print("暂无数据")
        exit(0)
    print("\n导出格式:  1.JSON  2.CSV  3.Excel  4.TXT  5.全部")
    choice = input("\n请选择 (1-5): ").strip()
    if choice == "1": export_json()
    elif choice == "2": export_csv()
    elif choice == "3": export_excel()
    elif choice == "4": export_txt()
    elif choice == "5":
        export_json(); export_csv(); export_excel(); export_txt()
    else: print("无效选择")
