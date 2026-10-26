"""
数据包导入模块 - 支持 CSV/XLSX 格式导入，智能列名映射

兼容大王解析、向风正在拼等工具导出的数据包格式。
"""

import csv
import io
import json
import re
import sqlite3
import uuid
from datetime import datetime
from pathlib import Path

from paths import DATA_DIR
from normalization import money, pdd_price, normalize_skus, normalize_goods, valid_goods_id
from storage import sync_skus
DB_PATH: Path = DATA_DIR / "goods.db"

COLUMN_ALIASES = {
    "goods_id": ["商品ID", "商品id", "goods_id", "id", "ID", "宝贝ID", "商品编号", "商品id号", "多多商品ID"],
    "title": ["商品标题", "标题", "title", "商品名称", "宝贝标题", "name", "商品名"],
    "subtitle": ["副标题", "卖点", "subtitle", "商品卖点", "短标题", "促销语"],
    "price": ["当前价格", "价格", "price", "售价", "券后价", "现价", "拼团价", "成交价", "团购价"],
    "original_price": ["原价", "市场价", "划线价", "original_price", "market_price", "吊牌价", "日常价"],
    "min_group_price": ["拼团价", "团购价", "group_price", "min_group_price", "最低拼团价"],
    "sales": ["销量", "sales", "已售", "月销", "总销量", "30天销量", "已售出"],
    "sold_quantity": ["已拼件数", "sold_quantity", "拼单数", "已拼团"],
    "shop_name": ["店铺名称", "店铺", "shop_name", "商家", "店铺名", "商家名称"],
    "shop_id": ["店铺ID", "shop_id", "商家ID", "店铺编号"],
    "main_images": ["主图链接", "主图", "main_images", "商品图片", "图片链接", "主图url", "商品主图"],
    "detail_images": ["详情图链接", "详情图", "detail_images", "详情图片", "详情页图片", "商品详情图"],
    "description": ["商品描述", "描述", "description", "详情", "商品详情", "宝贝描述"],
    "specs": ["规格信息", "规格", "specs", "商品规格", "属性规格", "规格名称"],
    "skus": ["SKU详情", "SKU", "skus", "规格明细", "SKU信息", "sku价格", "规格详情"],
    "attributes": ["商品属性", "属性", "attributes", "参数", "商品参数", "商品属性信息"],
    "source_url": ["商品链接", "链接", "source_url", "url", "商品URL", "宝贝链接", "商品地址"],
    "selling_points": ["卖点", "卖点标签", "标签", "selling_points", "商品卖点标签", "核心卖点"],
    "category": ["类目", "分类", "category", "商品类目", "所属分类", "商品分类"],
}

_IMAGE_SPLIT_RE = re.compile(r'[,;\s|]+')
_ATTR_SPLIT_RE = re.compile(r'[;；\n]+')
_ATTR_KV_RE = re.compile(r'^([^:：]+)[=:：]\s*(.+)$')
_SPEC_KV_RE = re.compile(r'^([^:：,，]+)[=:：]\s*(.+)$')
_PRICE_CLEAN_RE = re.compile(r'[¥￥$,，\s]')

_SKU_SEP_RE = re.compile(r'[|｜\n]+')
_SKU_FIELD_SEP = re.compile(r'[/,，;；\s]+')


def _detect_columns(headers: list) -> dict:
    detected = {}
    lower_headers = [h.strip().lower() for h in headers]
    for field, aliases in COLUMN_ALIASES.items():
        for idx, header in enumerate(lower_headers):
            if header in [a.lower() for a in aliases]:
                detected[field] = idx
                break
    return detected


def _parse_images(raw) -> list:
    if not raw:
        return []
    if isinstance(raw, list):
        return [u.strip() for u in raw if isinstance(u, str) and u.strip().startswith("http")]
    text = str(raw).strip()
    if not text:
        return []
    parts = _IMAGE_SPLIT_RE.split(text)
    result = []
    for p in parts:
        p = p.strip()
        if p.startswith("http") and p not in result:
            result.append(p)
    return result


def _parse_attributes(raw) -> dict:
    if not raw:
        return {}
    if isinstance(raw, dict):
        return {str(k): str(v) for k, v in raw.items() if k}
    text = str(raw).strip()
    if not text:
        return {}
    attrs = {}
    segments = _ATTR_SPLIT_RE.split(text)
    for seg in segments:
        seg = seg.strip()
        m = _ATTR_KV_RE.match(seg)
        if m:
            attrs[m.group(1).strip()] = m.group(2).strip()
        elif ":" in seg or "：" in seg:
            parts = re.split(r'[:：]', seg, maxsplit=1)
            if len(parts) == 2:
                attrs[parts[0].strip()] = parts[1].strip()
    return attrs


def _parse_specs(raw) -> dict:
    if not raw:
        return {}
    if isinstance(raw, dict):
        return raw
    text = str(raw).strip()
    if not text:
        return {}
    specs = {}
    segments = _ATTR_SPLIT_RE.split(text)
    for seg in segments:
        seg = seg.strip()
        m = _SPEC_KV_RE.match(seg)
        if m:
            name = m.group(1).strip()
            vals = m.group(2).strip()
            val_list = [v.strip() for v in re.split(r'[,，、]', vals) if v.strip()]
            if name and val_list:
                specs[name] = val_list
    return specs


def _parse_skus(raw, price_fallback: float = 0) -> list:
    if not raw:
        return []
    if isinstance(raw, list):
        result = []
        for item in raw:
            if isinstance(item, dict):
                result.append({
                    "sku_id": str(item.get("sku_id", "")),
                    "price": _to_float(item.get("price", price_fallback)),
                    "original_price": _to_float(item.get("original_price", 0)),
                    "stock": int(item["stock"]) if item.get("stock") is not None else None,
                    "image": item.get("image", item.get("thumb_url", "")),
                    "specs": _parse_sku_specs_from_dict(item.get("specs", {})),
                })
        return result
    text = str(raw).strip()
    if not text:
        return []
    skus = []
    parts = _SKU_SEP_RE.split(text)
    for part in parts:
        part = part.strip()
        if not part:
            continue
        fields = _SKU_FIELD_SEP.split(part)
        specs_dict = {}
        price = price_fallback
        stock = None
        image = ""
        spec_parts = []
        price_seen = False
        for f in fields:
            f = f.strip()
            if not f:
                continue
            if re.match(r'^[¥￥]?\d+\.?\d*$', f):
                if not price_seen:
                    price = money(f)
                    price_seen = True
                else:
                    stock = int(_to_float(f))
            elif re.match(r'^库存[:：]?\d+$', f):
                stock = int(re.search(r'\d+', f).group())
            elif f.startswith("http"):
                image = f
            else:
                spec_parts.append(f)
        if spec_parts:
            for sp in spec_parts:
                kv = re.split(r'[=:=：]', sp, maxsplit=1) if re.search(r'[=:=：]', sp) else [sp, sp]
                if len(kv) == 2:
                    specs_dict[kv[0].strip()] = kv[1].strip()
        skus.append({
            "sku_id": "",
            "price": price if price > 0 else price_fallback,
            "original_price": 0,
            "stock": stock,
            "image": image,
            "specs": specs_dict,
        })
    return skus


def _parse_sku_specs_from_dict(specs) -> dict:
    if isinstance(specs, dict):
        return {str(k): str(v) for k, v in specs.items() if k}
    if isinstance(specs, list):
        result = {}
        for item in specs:
            if isinstance(item, dict):
                k = item.get("spec_key", item.get("name", item.get("key", "")))
                v = item.get("spec_value", item.get("value", ""))
                if k and v:
                    result[str(k)] = str(v)
        return result
    return {}


def _to_float(value) -> float:
    try:
        if isinstance(value, str):
            value = _PRICE_CLEAN_RE.sub("", value)
        return float(value)
    except (ValueError, TypeError):
        return 0.0


def _to_int(value) -> int:
    try:
        return int(_to_float(value))
    except (ValueError, TypeError):
        return 0


def _generate_id(title: str) -> str:
    short = re.sub(r'[^\w]', '', title)[:10]
    return f"import_{short}_{uuid.uuid4().hex[:8]}"


def preview_file(filepath: str) -> dict:
    path = Path(filepath)
    suffix = path.suffix.lower()
    headers = []
    rows = []
    if suffix == ".csv":
        with open(path, "r", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            for i, row in enumerate(reader):
                if i == 0:
                    headers = [h.strip() for h in row]
                elif i <= 10:
                    rows.append(row)
                else:
                    break
    elif suffix in (".xlsx", ".xls"):
        try:
            import openpyxl
        except ImportError:
            return {"error": "需要安装 openpyxl: pip install openpyxl"}
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        ws = wb.active
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                headers = [str(h or "").strip() for h in row]
            elif i <= 10:
                rows.append([str(c or "") for c in row])
            else:
                break
        wb.close()
    else:
        return {"error": f"不支持的文件格式: {suffix}"}

    if not headers:
        return {"error": "未能读取到表头信息"}

    mapping = _detect_columns(headers)
    return {
        "headers": headers,
        "rows": rows,
        "detected_mapping": mapping,
        "mapping_labels": {k: headers[v] for k, v in mapping.items()},
    }


def import_file(filepath: str, column_mapping: dict = None) -> dict:
    path = Path(filepath)
    suffix = path.suffix.lower()
    all_rows = []
    headers = []
    if suffix == ".csv":
        with open(path, "r", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            for i, row in enumerate(reader):
                if i == 0:
                    headers = [h.strip() for h in row]
                else:
                    all_rows.append(row)
    elif suffix in (".xlsx", ".xls"):
        try:
            import openpyxl
        except ImportError:
            return {"error": "需要安装 openpyxl: pip install openpyxl"}
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        ws = wb.active
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                headers = [str(h or "").strip() for h in row]
            else:
                all_rows.append([str(c or "") for c in row])
        wb.close()
    else:
        return {"error": f"不支持的文件格式: {suffix}"}

    if not headers:
        return {"error": "未能读取到表头信息"}

    if column_mapping:
        mapping = column_mapping
    else:
        mapping = _detect_columns(headers)

    if "title" not in mapping and "goods_id" not in mapping:
        return {"error": "未能识别到标题或商品ID列，请手动指定列映射"}

    imported = 0
    skipped = 0
    errors = []
    conn = sqlite3.connect(DB_PATH)
    now = datetime.now().isoformat()

    for i, row in enumerate(all_rows):
        try:
            def _col(field):
                idx = mapping.get(field)
                if idx is not None and idx < len(row):
                    return row[idx].strip() if isinstance(row[idx], str) else str(row[idx] or "")
                return ""

            goods_id = _col("goods_id")
            title = _col("title")
            if not title and not goods_id:
                skipped += 1
                continue
            if not goods_id:
                goods_id = _generate_id(title)
            goods_id = valid_goods_id(goods_id)

            existing = conn.execute("SELECT 1 FROM goods WHERE goods_id = ?", (goods_id,)).fetchone()
            if existing:
                skipped += 1
                continue

            subtitle = _col("subtitle")
            price = money(_col("price"))
            original_price = money(_col("original_price"))
            min_group_price = money(_col("min_group_price"))
            if min_group_price <= 0:
                min_group_price = price
            sales = _to_int(_col("sales"))
            sold_quantity = _to_int(_col("sold_quantity"))
            shop_name = _col("shop_name")
            shop_id = _col("shop_id")
            description = _col("description")
            source_url = _col("source_url")

            main_images = _parse_images(_col("main_images"))
            detail_images = _parse_images(_col("detail_images"))
            specs = _parse_specs(_col("specs"))
            attributes = _parse_attributes(_col("attributes"))

            raw_sku = _col("skus")
            if raw_sku:
                skus = _parse_skus(raw_sku, price)
            else:
                skus = []
            skus = normalize_skus(skus)

            selling_points_raw = _col("selling_points")
            selling_points = []
            if selling_points_raw:
                selling_points = [s.strip() for s in re.split(r'[,，、]', selling_points_raw) if s.strip()]

            sku_images = []
            for sku in skus:
                img = sku.get("image", "")
                if img and img.startswith("http") and img not in sku_images:
                    sku_images.append(img)

            category = _col("category")
            raw_data = {}
            if category:
                raw_data["category"] = category

            conn.execute(
                """INSERT INTO goods
                   (goods_id, title, subtitle, selling_points, price, original_price,
                    min_group_price, sales, sold_quantity, shop_name, shop_id, shop_logo,
                    main_images, detail_images, sku_images, specs, skus, description,
                    attributes, source_url, crawl_time, updated_at, raw_data)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    goods_id, title, subtitle,
                    json.dumps(selling_points, ensure_ascii=False),
                    price, original_price, min_group_price,
                    sales, sold_quantity,
                    shop_name, shop_id, "",
                    json.dumps(main_images, ensure_ascii=False),
                    json.dumps(detail_images, ensure_ascii=False),
                    json.dumps(sku_images, ensure_ascii=False),
                    json.dumps(specs, ensure_ascii=False),
                    json.dumps(skus, ensure_ascii=False),
                    description,
                    json.dumps(attributes, ensure_ascii=False),
                    source_url, now, now,
                    json.dumps(raw_data, ensure_ascii=False) if raw_data else "{}",
                ),
            )
            _sync_sku_items(conn, goods_id, skus, now, supplied_fields=set(mapping))
            imported += 1
        except Exception as e:
            errors.append(f"第{i + 2}行: {str(e)}")

    conn.commit()
    conn.close()

    return {
        "total": len(all_rows),
        "imported": imported,
        "skipped": skipped,
        "errors": errors,
    }



def _extract_goods_from_raw_json(data: dict) -> dict | None:
    """从大王解析原始JSON格式中提取商品数据
    
    支持格式:
    - {"ts":..., "store": {"initDataObj": {"goods": {...}, "mall": {...}}}}
    - {"goods_id":..., ...} (直接格式)
    """
    if not data or not isinstance(data, dict):
        return None
    
    goods = None
    mall = None
    
    # 大王解析格式: store.initDataObj.goods
    if "store" in data:
        store = data.get("store", {})
        init = store.get("initDataObj", store.get("init_data_obj", {}))
        goods = init.get("goods", {})
        mall = init.get("mall")
    # 直接格式
    elif "goods_id" in data or "goodsID" in data:
        goods = data
    # goods_detail_response 格式
    elif "goods_detail_response" in data:
        resp = data["goods_detail_response"]
        gl = resp.get("goods", resp.get("goods_list", []))
        if isinstance(gl, list) and gl:
            goods = gl[0]
        elif isinstance(gl, dict):
            goods = gl
    
    if not goods:
        return None
    
    goods_id = str(goods.get("goods_id", goods.get("goodsID", "")))
    if not goods_id:
        return None
    goods_id = valid_goods_id(goods_id)
    
    # 提取SKU数据
    sku_list = goods.get("sku", goods.get("skus", []))
    specs = {}
    skus = []
    sku_images = []
    
    if isinstance(sku_list, list):
        for s in normalize_skus(sku_list, from_pdd=True):
            if not isinstance(s, dict):
                continue
            
            # 价格处理: 大王解析格式用 groupPrice/normalPrice (元), skuPrice (分)
            price = s["price"]
            orig_price = s["original_price"]
            stock = s["stock"]
            image = s.get("thumbUrl", s.get("thumb_url", s.get("image", "")))
            
            if image and image.startswith("http") and image not in sku_images:
                sku_images.append(image)
            
            sku_specs = []
            specs_dict = {}
            raw_specs = s.get("specs", s.get("spec", []))
            if isinstance(raw_specs, list):
                for sp in raw_specs:
                    if isinstance(sp, dict):
                        name = sp.get("key", sp.get("spec_key", sp.get("spec", "")))
                        value = sp.get("spec_value", sp.get("value", ""))
                        if name and value:
                            sku_specs.append({"key": name, "value": value})
                            specs_dict[name] = value
                            if name not in specs:
                                specs[name] = []
                            if value not in specs[name]:
                                specs[name].append(value)
            
            skus.append({
                "sku_id": str(s.get("skuID", s.get("skuId", s.get("sku_id", "")))),
                "specs": sku_specs,
                "price": price,
                "original_price": orig_price,
                "stock": stock,
                "image": image,
            })
    
    # 提取图片
    main_images = []
    for key in ["topGallery", "top_gallery", "gallery"]:
        gallery = goods.get(key, [])
        if isinstance(gallery, list):
            for item in gallery:
                if isinstance(item, str) and item.startswith("http"):
                    if item not in main_images:
                        main_images.append(item)
                elif isinstance(item, dict):
                    url = item.get("url", item.get("src", ""))
                    if url and url.startswith("http") and url not in main_images:
                        main_images.append(url)
    
    detail_images = []
    for key in ["detailGallery", "detail_gallery", "desc_gallery"]:
        gallery = goods.get(key, [])
        if isinstance(gallery, list):
            for item in gallery:
                if isinstance(item, str) and item.startswith("http"):
                    if item not in detail_images:
                        detail_images.append(item)
                elif isinstance(item, dict):
                    url = item.get("url", item.get("src", ""))
                    if url and url.startswith("http") and url not in detail_images:
                        detail_images.append(url)
    
    # 提取商品属性
    attributes = {}
    for key in ["goodsProperty", "goods_property", "goods_props", "attributes"]:
        props = goods.get(key, [])
        if isinstance(props, list):
            for p in props:
                if isinstance(p, dict):
                    name = p.get("key", p.get("name", ""))
                    values = p.get("values", [])
                    if name and values:
                        attributes[name] = str(values[0]) if isinstance(values, list) and values else str(values)
                    elif name and p.get("value"):
                        attributes[name] = str(p["value"])
    
    # 提取卖点/服务标签
    selling_points = []
    for key in ["servicePromise", "service_promise", "service_tags", "goods_service_tags"]:
        tags = goods.get(key, [])
        if isinstance(tags, list):
            for tag in tags:
                if isinstance(tag, dict):
                    text = tag.get("type", tag.get("text", tag.get("name", tag.get("desc", ""))))
                    if text and text not in selling_points:
                        selling_points.append(text)
    
    # 价格处理
    price = pdd_price(goods)
    if price <= 0 and skus:
        prices = [s["price"] for s in skus if s["price"] > 0]
        if prices:
            price = min(prices)
    
    orig_price = pdd_price(goods, original=True)
    
    result = {
        "goods_id": goods_id,
        "title": goods.get("goodsName", goods.get("goods_name", goods.get("title", ""))),
        "subtitle": goods.get("goods_desc", goods.get("desc", "")),
        "price": price,
        "original_price": orig_price,
        "min_group_price": price,
        "sales": int(goods.get("sales", goods.get("sold_quantity", 0)) or 0),
        "sold_quantity": int(goods.get("sold_quantity", goods.get("sales", 0)) or 0),
        "shop_name": (mall or {}).get("mallName", (mall or {}).get("mall_name", goods.get("mall_name", ""))),
        "shop_id": str((mall or {}).get("mallID", (mall or {}).get("mall_id", goods.get("mall_id", "")))),
        "shop_logo": (mall or {}).get("mallLogo", (mall or {}).get("mall_logo", goods.get("mall_logo", ""))),
        "main_images": main_images,
        "detail_images": detail_images,
        "sku_images": sku_images,
        "specs": specs,
        "skus": skus,
        "description": goods.get("goods_desc", goods.get("desc", "")),
        "attributes": attributes,
        "selling_points": selling_points,
        "source_url": f"https://mobile.yangkeduo.com/goods.html?goods_id={goods_id}",
    }
    result["raw_data"] = dict(result)
    return result


def import_raw_json_files(file_paths: list) -> dict:
    """导入大王解析原始JSON文件到数据库"""
    conn = sqlite3.connect(DB_PATH)
    now = datetime.now().isoformat()
    imported = 0
    skipped = 0
    errors = []
    
    for filepath in file_paths:
        try:
            path = Path(filepath)
            with open(path, "r", encoding="utf-8") as f:
                raw_data = json.load(f)
            
            goods_info = _extract_goods_from_raw_json(raw_data)
            if not goods_info:
                errors.append(f"{path.name}: 无法提取商品数据")
                continue
            
            goods_id = goods_info["goods_id"]
            existing = conn.execute("SELECT 1 FROM goods WHERE goods_id = ?", (goods_id,)).fetchone()
            if existing:
                skipped += 1
                continue
            
            conn.execute(
                """INSERT INTO goods
                   (goods_id, title, subtitle, selling_points, price, original_price,
                    min_group_price, sales, sold_quantity, shop_name, shop_id, shop_logo,
                    main_images, detail_images, sku_images, specs, skus, description,
                    attributes, source_url, crawl_time, updated_at, raw_data)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    goods_id, goods_info["title"], goods_info["subtitle"],
                    json.dumps(goods_info["selling_points"], ensure_ascii=False),
                    goods_info["price"], goods_info["original_price"], goods_info["min_group_price"],
                    goods_info["sales"], goods_info["sold_quantity"],
                    goods_info["shop_name"], goods_info["shop_id"], goods_info["shop_logo"],
                    json.dumps(goods_info["main_images"], ensure_ascii=False),
                    json.dumps(goods_info["detail_images"], ensure_ascii=False),
                    json.dumps(goods_info["sku_images"], ensure_ascii=False),
                    json.dumps(goods_info["specs"], ensure_ascii=False),
                    json.dumps(goods_info["skus"], ensure_ascii=False),
                    goods_info["description"],
                    json.dumps(goods_info["attributes"], ensure_ascii=False),
                    goods_info["source_url"], now, now,
                    json.dumps(goods_info["raw_data"], ensure_ascii=False),
                ),
            )
            _sync_sku_items(conn, goods_id, goods_info["skus"], now)
            imported += 1
        except Exception as e:
            errors.append(f"{Path(filepath).name}: {str(e)}")
    
    conn.commit()
    conn.close()
    return {"imported": imported, "skipped": skipped, "errors": errors}


def import_raw_json_directory(dir_path: str) -> dict:
    """导入目录中所有大王解析原始JSON文件"""
    path = Path(dir_path)
    if not path.exists():
        return {"error": f"目录不存在: {dir_path}"}
    
    json_files = []
    for f in path.rglob("*.txt"):
        try:
            with open(f, "r", encoding="utf-8") as fh:
                content = fh.read(200)
                if content.strip().startswith("{"):
                    json_files.append(str(f))
        except:
            pass
    for f in path.rglob("*.json"):
        json_files.append(str(f))
    
    if not json_files:
        return {"error": "未找到JSON格式的文件"}
    
    return import_raw_json_files(json_files)

# ---- import_goods_json: 本系统 JSON 格式导入 ----

def import_goods_json(file_path: str) -> dict:
    """导入本系统导出的JSON格式文件，支持更新已有商品。
    修正：同步更新 price_min/price_max/sku_items/price_modified。"""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    now = datetime.now().isoformat()
    imported = 0
    updated = 0
    errors = []

    try:
        with open(file_path, "r", encoding="utf-8") as f:
            goods_list = json.load(f)

        if isinstance(goods_list, dict):
            goods_list = [goods_list]

        for goods in goods_list:
            goods_id = goods.get("goods_id", "") if isinstance(goods, dict) else ""
            if not goods_id:
                continue

            try:
                provided_fields = set(goods)
                goods = normalize_goods({**goods, "capture_source": "import"})
                existing = conn.execute(
                    "SELECT 1 FROM goods WHERE goods_id = ?", (goods_id,)
                ).fetchone()

                price = goods.get("price", 0) or 0
                original_price = goods.get("original_price", 0) or 0
                min_group_price = goods.get("min_group_price", 0) or 0
                title = goods.get("title", "")
                subtitle = goods.get("subtitle", "")
                selling_points = goods.get("selling_points", [])
                sales = goods.get("sales", 0) or 0
                sold_quantity = goods.get("sold_quantity", 0) or 0
                shop_name = goods.get("shop_name", "")
                shop_id = goods.get("shop_id", "")
                main_images = goods.get("main_images", [])
                detail_images = goods.get("detail_images", [])
                sku_images = goods.get("sku_images", [])
                specs = goods.get("specs", {})
                skus = goods.get("skus", [])
                description = goods.get("description", "")
                attributes = goods.get("attributes", {})
                source_url = goods.get("source_url", "")
                crawl_time = goods.get("crawl_time", now)

                # 计算 price_min / price_max
                all_prices = [
                    s.get("price", 0)
                    for s in skus
                    if isinstance(s, dict) and (s.get("price", 0) or 0) > 0
                ]
                if not all_prices and price > 0:
                    all_prices = [price]
                price_min = min(all_prices) if all_prices else 0
                price_max = max(all_prices) if all_prices else 0

                if existing:
                    conn.execute(
                        """UPDATE goods SET
                           title=?, subtitle=?, selling_points=?, price=?, original_price=?,
                           min_group_price=?, sales=?, sold_quantity=?, shop_name=?, shop_id=?,
                           main_images=?, detail_images=?, sku_images=?, specs=?, skus=?,
                           description=?, attributes=?, source_url=?,
                           price_min=?, price_max=?, price_modified=1, updated_at=?
                           WHERE goods_id=?""",
                        (
                            title, subtitle,
                            json.dumps(selling_points, ensure_ascii=False),
                            price, original_price, min_group_price,
                            sales, sold_quantity, shop_name, shop_id,
                            json.dumps(main_images, ensure_ascii=False),
                            json.dumps(detail_images, ensure_ascii=False),
                            json.dumps(sku_images, ensure_ascii=False),
                            json.dumps(specs, ensure_ascii=False),
                            json.dumps(skus, ensure_ascii=False),
                            description,
                            json.dumps(attributes, ensure_ascii=False),
                            source_url,
                            price_min, price_max, now,
                            goods_id,
                        ),
                    )
                    # 同步 sku_items 表
                    _sync_sku_items(conn, goods_id, skus, now, supplied_fields=provided_fields)
                    updated += 1
                else:
                    conn.execute(
                        """INSERT INTO goods
                           (goods_id, title, subtitle, selling_points, price, original_price,
                            min_group_price, sales, sold_quantity, shop_name, shop_id, shop_logo,
                            main_images, detail_images, sku_images, specs, skus, description,
                            attributes, source_url, crawl_time, updated_at, raw_data,
                            price_min, price_max, price_modified)
                           VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                        (
                            goods_id, title, subtitle,
                            json.dumps(selling_points, ensure_ascii=False),
                            price, original_price, min_group_price,
                            sales, sold_quantity, shop_name, shop_id, "",
                            json.dumps(main_images, ensure_ascii=False),
                            json.dumps(detail_images, ensure_ascii=False),
                            json.dumps(sku_images, ensure_ascii=False),
                            json.dumps(specs, ensure_ascii=False),
                            json.dumps(skus, ensure_ascii=False),
                            description,
                            json.dumps(attributes, ensure_ascii=False),
                            source_url, crawl_time, now, "{}",
                            price_min, price_max, 1,
                        ),
                    )
                    _sync_sku_items(conn, goods_id, skus, now, supplied_fields=provided_fields)
                    imported += 1
            except Exception as e:
                errors.append(f"{goods_id}: {str(e)}")
    except Exception as e:
        errors.append(f"文件解析失败: {str(e)}")

    conn.commit()
    conn.close()
    return {"imported": imported, "updated": updated, "skipped": 0, "errors": errors}


def _sync_sku_items(conn, goods_id: str, skus: list, now: str, *, supplied_fields=None):
    skus = normalize_skus(skus)
    sync_skus(conn, goods_id, skus, now)
    prices = [s["price"] for s in skus if s["price"] > 0]
    if not prices:
        row = conn.execute("SELECT price FROM goods WHERE goods_id=?", (goods_id,)).fetchone()
        prices = [row[0] or 0] if row else [0]
    stocks = [s["stock"] for s in skus]
    total = sum(stocks) if stocks and all(s is not None for s in stocks) else None
    conn.execute("UPDATE goods SET skus=?, price_min=?, price_max=?, total_stock=? WHERE goods_id=?",
                 (json.dumps(skus, ensure_ascii=False), min(prices), max(prices), total, goods_id))
    from provenance import mark_import
    mark_import(conn, goods_id, now, supplied_fields)



if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("用法: python importer.py <文件路径.csv|xlsx>")
        sys.exit(1)
    filepath = sys.argv[1]
    print("预览文件...")
    result = preview_file(filepath)
    if "error" in result:
        print(f"错误: {result['error']}")
        sys.exit(1)
    print(f"表头: {result['headers']}")
    print(f"检测到的映射: {result['mapping_labels']}")
    print(f"前{len(result['rows'])}行数据:")
    for row in result["rows"]:
        print(f"  {row[:5]}...")
    confirm = input("\n确认导入? (y/n): ").strip().lower()
    if confirm == "y":
        result = import_file(filepath)
        print(f"导入完成: {result['imported']} 条成功, {result['skipped']} 条跳过, {len(result['errors'])} 条失败")
        if result["errors"]:
            for e in result["errors"][:10]:
                print(f"  - {e}")
